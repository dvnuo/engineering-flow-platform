"""PowerPoint (PPTX) parser: slide text, tables and speaker notes as markdown."""

from __future__ import annotations

import time
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List

from .docx import _table_to_json, _table_to_markdown
from .models import Block, ParseResult

PPTX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _iter_shapes(shapes) -> Iterable:
    """Walk shapes depth first so text inside grouped shapes is not lost."""
    for shape in shapes:
        yield shape
        nested = getattr(shape, "shapes", None)
        if nested is not None and getattr(shape, "shape_type", None) is not None:
            try:
                from pptx.enum.shapes import MSO_SHAPE_TYPE

                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from _iter_shapes(nested)
            except Exception:
                continue


def _shape_paragraphs(shape) -> List[str]:
    if not getattr(shape, "has_text_frame", False):
        return []
    text_frame = shape.text_frame
    if text_frame is None:
        return []
    lines: List[str] = []
    for paragraph in text_frame.paragraphs:
        text = "".join(run.text for run in paragraph.runs).strip() or paragraph.text.strip()
        if text:
            lines.append(text)
    return lines


def _shape_table_rows(shape) -> List[List[str]]:
    if not getattr(shape, "has_table", False):
        return []
    try:
        table = shape.table
    except Exception:
        return []
    rows: List[List[str]] = []
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    return [row for row in rows if any(cell for cell in row)]


def _notes_text(slide) -> str:
    try:
        if not slide.has_notes_slide:
            return ""
        frame = slide.notes_slide.notes_text_frame
    except Exception:
        return ""
    if frame is None:
        return ""
    return frame.text.strip()


async def parse_pptx(file_path: str, options: Dict = None) -> ParseResult:
    """Parse a PPTX deck into one heading block per slide plus its text, tables and notes."""
    start_time = time.time()
    options = options or {}
    max_slides = int(options.get("max_slides", 500) or 500)

    file_id = Path(file_path).stem.split("_")[0]
    filename = Path(file_path).name

    try:
        from pptx import Presentation
    except ImportError:
        return ParseResult(
            success=False,
            content_type=PPTX_MIME_TYPE,
            file_id=file_id,
            filename=filename,
            error="python-pptx not installed",
        )

    try:
        presentation = Presentation(file_path)
        blocks: List[Block] = []
        extracted_at = datetime.now().isoformat()
        text_shapes = 0
        tables = 0
        slide_count = 0

        for slide_number, slide in enumerate(presentation.slides, 1):
            if slide_number > max_slides:
                break
            slide_count += 1
            title_shape = None
            try:
                title_shape = slide.shapes.title
            except Exception:
                title_shape = None
            title_lines = _shape_paragraphs(title_shape) if title_shape is not None else []
            title = " ".join(title_lines).strip()
            title_shape_id = getattr(title_shape, "shape_id", None)

            blocks.append(
                Block(
                    chunk_id=f"{file_id}_pptx_{slide_number}_0",
                    type="heading",
                    content=f"Slide {slide_number}: {title}" if title else f"Slide {slide_number}",
                    level=2,
                    page=slide_number,
                    method="python-pptx",
                    confidence=0.95,
                    extracted_at=extracted_at,
                )
            )

            index = 1
            for shape in _iter_shapes(slide.shapes):
                if title_shape_id is not None and getattr(shape, "shape_id", None) == title_shape_id:
                    continue
                rows = _shape_table_rows(shape)
                if rows:
                    tables += 1
                    blocks.append(
                        Block(
                            chunk_id=f"{file_id}_pptx_{slide_number}_{index}",
                            type="table",
                            content="",
                            markdown=_table_to_markdown(rows),
                            table_json=_table_to_json(rows),
                            page=slide_number,
                            method="python-pptx",
                            confidence=0.95,
                            extracted_at=extracted_at,
                        )
                    )
                    index += 1
                    continue
                lines = _shape_paragraphs(shape)
                if not lines:
                    continue
                text_shapes += 1
                blocks.append(
                    Block(
                        chunk_id=f"{file_id}_pptx_{slide_number}_{index}",
                        type="paragraph",
                        content="\n".join(lines),
                        page=slide_number,
                        method="python-pptx",
                        confidence=0.95,
                        extracted_at=extracted_at,
                    )
                )
                index += 1

            notes = _notes_text(slide)
            if notes:
                blocks.append(
                    Block(
                        chunk_id=f"{file_id}_pptx_{slide_number}_{index}",
                        type="paragraph",
                        content=f"Speaker notes: {notes}",
                        page=slide_number,
                        method="python-pptx",
                        confidence=0.9,
                        extracted_at=extracted_at,
                    )
                )

        # A deck whose slides carry no text at all (pictures only) has nothing
        # for the model; report that instead of an empty success.
        has_text = any(
            (block.type != "heading" and (block.content or block.markdown))
            for block in blocks
        )
        if not has_text:
            return ParseResult(
                success=False,
                content_type=PPTX_MIME_TYPE,
                file_id=file_id,
                filename=filename,
                error="The presentation contains no text, tables or notes to extract",
                parse_time_ms=int((time.time() - start_time) * 1000),
            )

        return ParseResult(
            success=True,
            content_type=PPTX_MIME_TYPE,
            file_id=file_id,
            filename=filename,
            markdown=_blocks_to_markdown(blocks),
            blocks=blocks,
            json={"slides": slide_count, "text_shapes": text_shapes, "tables": tables},
            parse_time_ms=int((time.time() - start_time) * 1000),
        )
    except Exception as exc:
        return ParseResult(
            success=False,
            content_type=PPTX_MIME_TYPE,
            file_id=file_id,
            filename=filename,
            error=str(exc),
            parse_time_ms=int((time.time() - start_time) * 1000),
        )


def _blocks_to_markdown(blocks: List[Block]) -> str:
    parts: List[str] = []
    for block in blocks:
        if block.type == "heading":
            parts.append(f"\n## {block.content}\n")
        elif block.type == "table" and block.markdown:
            parts.append(f"\n{block.markdown}\n")
        elif block.type == "paragraph":
            if block.content.startswith("Speaker notes: "):
                parts.append(f"> {block.content}\n")
            else:
                parts.append(f"{block.content}\n")
    return "\n".join(parts).strip() + "\n"
