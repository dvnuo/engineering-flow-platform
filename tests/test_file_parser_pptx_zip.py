"""PowerPoint and ZIP projections for chat attachments."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest

from src.utils.file_parser import archive as archive_module
from src.utils.file_parser.archive import parse_zip
from src.utils.file_parser.pptx import parse_pptx


def _write_pptx(path: Path, *, with_text: bool = True) -> None:
    from pptx import Presentation

    prs = Presentation()
    if with_text:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Roadmap"
        body = slide.placeholders[1].text_frame
        body.text = "Ship uploads"
        body.add_paragraph().text = "Ship pptx parsing"
        slide.notes_slide.notes_text_frame.text = "Keep it short"

        second = prs.slides.add_slide(prs.slide_layouts[5])
        second.shapes.title.text = "Budget"
        table = second.shapes.add_table(3, 2, 0, 0, 100, 100).table
        table.cell(0, 0).text = "Item"
        table.cell(0, 1).text = "Cost"
        table.cell(1, 0).text = "Hosting"
        table.cell(1, 1).text = "200"
        table.cell(2, 0).text = "Licences"
        table.cell(2, 1).text = "50"
    else:
        prs.slides.add_slide(prs.slide_layouts[6])  # blank layout, nothing on it
    prs.save(str(path))


@pytest.mark.asyncio
async def test_parse_pptx_projects_titles_bullets_tables_and_notes(tmp_path: Path):
    deck = tmp_path / "abc123_deck.pptx"
    _write_pptx(deck)

    result = await parse_pptx(str(deck))

    assert result.success is True, result.error
    assert result.file_id == "abc123"
    assert result.json == {"slides": 2, "text_shapes": 1, "tables": 1}
    assert "## Slide 1: Roadmap" in result.markdown
    assert "Ship uploads\nShip pptx parsing" in result.markdown
    assert "> Speaker notes: Keep it short" in result.markdown
    assert "## Slide 2: Budget" in result.markdown
    assert "| Item" in result.markdown and "| Hosting" in result.markdown

    headings = [b for b in result.blocks if b.type == "heading"]
    assert [b.page for b in headings] == [1, 2]
    assert [b.content for b in headings] == ["Slide 1: Roadmap", "Slide 2: Budget"]
    tables = [b for b in result.blocks if b.type == "table"]
    assert len(tables) == 1 and tables[0].page == 2
    assert tables[0].table_json[0] == ["Item", "Cost"]
    # The title text is not duplicated as a paragraph.
    assert not any(b.type == "paragraph" and b.content == "Roadmap" for b in result.blocks)


@pytest.mark.asyncio
async def test_parse_pptx_reports_a_deck_without_text(tmp_path: Path):
    deck = tmp_path / "empty_deck.pptx"
    _write_pptx(deck, with_text=False)

    result = await parse_pptx(str(deck))

    assert result.success is False
    assert "no text" in (result.error or "")


def _write_zip(path: Path, entries: dict) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries.items():
            zf.writestr(name, data)


@pytest.mark.asyncio
async def test_parse_zip_lists_everything_and_inlines_text_members(tmp_path: Path):
    archive = tmp_path / "f1_bundle.zip"
    _write_zip(
        archive,
        {
            "app/README.md": "# App\n\nFirst paragraph.\n\nSecond paragraph.\n",
            "app/settings.yaml": "debug: false\n",
            "app/Dockerfile": "FROM python:3.12\n",
            "app/logo.png": b"\x89PNG\r\n\x1a\n\x00\x00binary",
            "app/vendor.zip": b"PK\x03\x04",
            "__MACOSX/._README.md": b"junk",
            "app/.DS_Store": b"junk",
        },
    )

    result = await parse_zip(str(archive))

    assert result.success is True
    assert result.file_id == "f1"
    assert result.json["entries"] == 5
    assert result.json["text_files"] == 3
    assert {item["name"]: item["reason"] for item in result.json["skipped"]} == {
        "app/logo.png": "binary",
        "app/vendor.zip": "nested archive",
    }
    md = result.markdown
    assert md.startswith("# Archive: f1_bundle.zip")
    assert "Archive f1_bundle.zip: 5 file(s)" in md
    assert "- app/logo.png (" in md
    assert "__MACOSX" not in md and ".DS_Store" not in md
    assert "## app/README.md" in md and "Second paragraph." in md
    assert "## app/Dockerfile" in md and "FROM python:3.12" in md
    assert "Not inlined:" in md

    listing = [b for b in result.blocks if b.type == "list"]
    assert len(listing) == 1 and listing[0].page == 1
    readme_chunks = [b for b in result.blocks if b.sheet == "app/README.md"]
    assert [b.type for b in readme_chunks] == ["heading", "paragraph", "paragraph", "paragraph"]
    assert all(b.page == readme_chunks[0].page for b in readme_chunks)


@pytest.mark.asyncio
async def test_parse_zip_applies_size_and_text_budgets(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(archive_module, "MAX_ENTRY_BYTES", 64)
    archive = tmp_path / "f2_big.zip"
    _write_zip(
        archive,
        {
            "small.txt": "tiny",
            "large.txt": "x" * 200,  # over MAX_ENTRY_BYTES: listed, not inlined
            "notes.md": "word " * 10,  # 50 bytes: inlined, then cut by max_file_chars
        },
    )

    result = await parse_zip(str(archive), {"display_name": "big.zip", "max_file_chars": 20, "max_total_chars": 30})

    assert result.success is True
    assert result.filename == "big.zip"
    assert "# Archive: big.zip" in result.markdown
    skipped = {item["name"]: item["reason"] for item in result.json["skipped"]}
    assert skipped["large.txt"] == "too large to inline"
    assert result.json["text_files"] == 2
    assert result.json["total_chars"] <= 30
    assert "## notes.md (truncated)" in result.markdown
    assert "## small.txt" in result.markdown


@pytest.mark.asyncio
async def test_parse_zip_decodes_gbk_members(tmp_path: Path):
    archive = tmp_path / "f5_gbk.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("导出/说明.txt", "第三季度营收增长 12%".encode("gb18030"))

    result = await parse_zip(str(archive))

    assert result.success is True
    assert result.json["text_files"] == 1
    assert "第三季度营收增长 12%" in result.markdown


@pytest.mark.asyncio
async def test_parse_text_file_decodes_gbk(tmp_path: Path):
    from src.utils.file_parser.text import parse_text_file

    log = tmp_path / "f6.log"
    log.write_bytes("2026-09-16 错误：连接超时\r\n第二行\r\n".encode("gb18030"))

    result = await parse_text_file(str(log), file_id="f6", filename="app.log", content_type="text/plain")

    assert result.success is True
    assert result.markdown == "2026-09-16 错误：连接超时\n第二行\n"
    assert result.json == {"encoding": "gb18030", "chars": len(result.markdown)}


@pytest.mark.asyncio
async def test_parse_zip_rejects_a_non_archive(tmp_path: Path):
    bogus = tmp_path / "f3_bogus.zip"
    bogus.write_bytes(b"not a zip")

    result = await parse_zip(str(bogus))

    assert result.success is False
    assert "Not a valid zip archive" in (result.error or "")


@pytest.mark.asyncio
async def test_parse_zip_skips_suspiciously_compressed_members(tmp_path: Path, monkeypatch):
    archive = tmp_path / "f4_bomb.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("zeros.txt", b"0" * (2 * 1024 * 1024))  # compresses ~1000x
        zf.writestr("ok.txt", "fine")

    result = await parse_zip(str(archive))

    assert result.success is True
    skipped = {item["name"]: item["reason"] for item in result.json["skipped"]}
    assert skipped["zeros.txt"] == "too large to inline"
    assert "## ok.txt" in result.markdown
