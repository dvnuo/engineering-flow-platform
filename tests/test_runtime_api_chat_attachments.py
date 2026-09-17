"""Chat attachment API used by the Portal chatbox.

The Runtime v2 rewrite (#521) dropped ``/api/files/upload`` and friends while
the Portal kept proxying chatbox uploads to them, so every "attach a file and
ask the model" attempt failed with 502. These tests pin the restored routes end
to end: upload -> parse -> the chat handlers can consume the attachment.
"""

from __future__ import annotations

import sys

import pytest
from aiohttp import FormData, web
from aiohttp.test_utils import TestClient, TestServer


PNG_BYTES = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01"
    b"\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)


@pytest.fixture
def isolated_attachment_storage(tmp_path, monkeypatch):
    """Point upload + file-context storage at tmp_path and clear env overrides.

    Other test modules load ``src.utils.file_parser`` through lightweight
    loaders that drop the real package from ``sys.modules`` afterwards, so a
    fresh ``import`` here could hand back a different module object than the
    one ``runtime_api`` bound at import time. Patch the namespaces the runtime
    handlers actually use instead.
    """
    from src.gateway import runtime_api

    storage_namespace = runtime_api.get_metadata.__globals__
    # The package namespace runtime_api's upload_file lives in holds its
    # storage submodule; register it so the handlers' local
    # ``from src.utils.file_parser.storage import ...`` resolve to the same
    # module whose globals are patched below.
    storage_module = runtime_api.upload_file.__globals__["storage"]
    assert storage_module.__dict__ is storage_namespace
    monkeypatch.setitem(sys.modules, "src.utils.file_parser.storage", storage_module)
    fc_storage = runtime_api.file_context_storage

    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    monkeypatch.setitem(storage_namespace, "UPLOAD_DIR", upload_dir)
    monkeypatch.setitem(storage_namespace, "METADATA_FILE", upload_dir / "metadata.json")
    monkeypatch.setitem(storage_namespace, "_file_metadata", {})

    ctx_dir = tmp_path / "file_context"
    (ctx_dir / "sessions").mkdir(parents=True)
    (ctx_dir / "chunks").mkdir(parents=True)
    monkeypatch.setattr(fc_storage, "base_dir", ctx_dir)
    monkeypatch.setattr(fc_storage, "sessions_dir", ctx_dir / "sessions")
    monkeypatch.setattr(fc_storage, "chunks_dir", ctx_dir / "chunks")

    monkeypatch.delenv("EFP_CHAT_UPLOAD_EXTENSIONS", raising=False)
    monkeypatch.delenv("EFP_MAX_UPLOAD_MB", raising=False)
    return upload_dir


async def _client(client_max_size: int | None = None) -> TestClient:
    from src.gateway.runtime_api import setup_runtime_api_routes

    kwargs = {"client_max_size": client_max_size} if client_max_size else {}
    app = web.Application(**kwargs)
    setup_runtime_api_routes(app)
    client = TestClient(TestServer(app))
    await client.start_server()
    return client


def _file_form(filename: str, data: bytes, content_type: str = "application/octet-stream") -> FormData:
    # Browsers send the multipart filename as raw UTF-8; aiohttp's client would
    # percent-encode it unless told not to, which no browser does.
    form = FormData(quote_fields=False)
    form.add_field("file", data, filename=filename, content_type=content_type)
    return form


@pytest.mark.asyncio
async def test_upload_then_parse_feeds_the_chat_attachment_context(isolated_attachment_storage):
    from src.gateway import runtime_api

    fc_storage = runtime_api.file_context_storage
    client = await _client()
    try:
        upload = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("notes.txt", b"hello world\n\nsecond paragraph\n", "text/plain"),
        )
        assert upload.status == 201, await upload.text()
        uploaded = await upload.json()
        assert uploaded["success"] is True
        assert uploaded["filename"] == "notes.txt"
        assert uploaded["content_type"] == "text/plain"
        assert uploaded["size"] == len(b"hello world\n\nsecond paragraph\n")
        assert uploaded["session_id"] == "s1"
        file_id = uploaded["file_id"]
        assert file_id

        parsed = await client.post(
            "/api/files/parse",
            params={"session_id": "s1"},
            json={"file_id": file_id, "options": {}},
        )
        assert parsed.status == 200, await parsed.text()
        body = await parsed.json()
        assert body["success"] is True
        assert body["file_id"] == file_id
        assert "hello world" in body["markdown"]
        assert body["saved_chunks"] >= 1
        assert body["blocks"]
    finally:
        await client.close()

    meta = fc_storage.get_file_meta("s1", file_id)
    assert meta is not None
    assert meta.parse_status == "completed"

    # The chat handlers resolve attachments through the same storage, so a
    # parsed upload is ready to be sent to the model as file context.
    context = await runtime_api._ensure_chat_attachment_context(session_id="s1", attachment_ids=[file_id])
    assert context["context_file_ids"] == [file_id]
    assert context["already_ready_file_ids"] == [file_id]
    assert context["failures"] == []


@pytest.mark.asyncio
async def test_unparsed_upload_is_parsed_on_demand_by_the_chat_handlers(isolated_attachment_storage):
    from src.gateway import runtime_api

    client = await _client()
    try:
        upload = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("data.csv", b"name,age\nalice,30\nbob,41\n", "text/csv"),
        )
        assert upload.status == 201, await upload.text()
        file_id = (await upload.json())["file_id"]
    finally:
        await client.close()

    context = await runtime_api._ensure_chat_attachment_context(session_id="s1", attachment_ids=[file_id])
    assert context["context_file_ids"] == [file_id]
    assert context["parsed_file_ids"] == [file_id]
    assert context["failures"] == []


@pytest.mark.asyncio
async def test_image_upload_is_sent_to_the_model_as_an_image(isolated_attachment_storage, monkeypatch):
    from src.gateway import runtime_api

    # A deployment whose model can see adds image types to the allowlist.
    monkeypatch.setenv("EFP_CHAT_UPLOAD_EXTENSIONS", "png,jpg,pdf")
    client = await _client()
    try:
        upload = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("shot.png", PNG_BYTES, "image/png"),
        )
        assert upload.status == 201, await upload.text()
        uploaded = await upload.json()
        assert uploaded["content_type"] == "image/png"
        file_id = uploaded["file_id"]

        raw = await client.get(f"/api/files/{file_id}", params={"session_id": "s1"})
        assert raw.status == 200
        assert raw.headers["Content-Type"].startswith("image/png")
        assert await raw.read() == PNG_BYTES
    finally:
        await client.close()

    context = await runtime_api._ensure_chat_attachment_context(session_id="s1", attachment_ids=[file_id])
    assert context["image_file_ids"] == [file_id]
    assert context["context_file_ids"] == []
    images = await runtime_api._collect_attached_images(session_id="s1", message="", attachments=[file_id])
    assert len(images) == 1
    assert images[0].startswith("data:image/png;base64,")


def _sample_pptx_bytes() -> bytes:
    import io

    from pptx import Presentation

    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[1])
    title_slide.shapes.title.text = "Quarterly review"
    title_slide.placeholders[1].text_frame.text = "Revenue grew 12%"
    title_slide.placeholders[1].text_frame.add_paragraph().text = "Churn fell to 3%"
    title_slide.notes_slide.notes_text_frame.text = "Mention the onboarding flow"

    table_slide = prs.slides.add_slide(prs.slide_layouts[5])
    table_slide.shapes.title.text = "Numbers"
    shape = table_slide.shapes.add_table(2, 2, 0, 0, 100, 100)
    shape.table.cell(0, 0).text = "Region"
    shape.table.cell(0, 1).text = "Revenue"
    shape.table.cell(1, 0).text = "EMEA"
    shape.table.cell(1, 1).text = "1.2M"

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _sample_zip_bytes() -> bytes:
    import io
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("project/README.md", "# Demo project\n\nRuns the nightly export.\n")
        zf.writestr("project/config.json", '{"retries": 3}')
        zf.writestr("project/bin/tool.exe", b"MZ\x90\x00\x00\x00binary")
        zf.writestr("project/vendor.zip", b"PK\x03\x04nested")
        zf.writestr("__MACOSX/._README.md", b"junk")
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_pptx_upload_is_projected_to_slide_text(isolated_attachment_storage):
    from src.gateway import runtime_api

    client = await _client()
    try:
        upload = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("review.pptx", _sample_pptx_bytes(), "application/octet-stream"),
        )
        assert upload.status == 201, await upload.text()
        uploaded = await upload.json()
        assert uploaded["content_type"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        file_id = uploaded["file_id"]

        parsed = await client.post("/api/files/parse", params={"session_id": "s1"}, json={"file_id": file_id})
        assert parsed.status == 200, await parsed.text()
        body = await parsed.json()
        assert body["success"] is True
        markdown = body["markdown"]
        assert "## Slide 1: Quarterly review" in markdown
        assert "Revenue grew 12%" in markdown
        assert "Churn fell to 3%" in markdown
        assert "Speaker notes: Mention the onboarding flow" in markdown
        assert "## Slide 2: Numbers" in markdown
        assert "| EMEA" in markdown
        assert body["json"]["slides"] == 2
        assert body["json"]["tables"] == 1
    finally:
        await client.close()

    context = await runtime_api._ensure_chat_attachment_context(session_id="s1", attachment_ids=[file_id])
    assert context["context_file_ids"] == [file_id]
    assert context["failures"] == []


@pytest.mark.asyncio
async def test_zip_upload_is_projected_to_listing_and_text_members(isolated_attachment_storage):
    from src.gateway import runtime_api

    client = await _client()
    try:
        upload = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("project.zip", _sample_zip_bytes(), "application/zip"),
        )
        assert upload.status == 201, await upload.text()
        uploaded = await upload.json()
        assert uploaded["content_type"] == "application/zip"
        file_id = uploaded["file_id"]

        parsed = await client.post("/api/files/parse", params={"session_id": "s1"}, json={"file_id": file_id})
        assert parsed.status == 200, await parsed.text()
        body = await parsed.json()
        assert body["success"] is True
        markdown = body["markdown"]
        assert "# Archive: project.zip" in markdown
        assert "- project/README.md" in markdown
        assert "- project/bin/tool.exe" in markdown
        assert "__MACOSX" not in markdown
        assert "## project/README.md" in markdown
        assert "Runs the nightly export." in markdown
        assert '"retries": 3' in markdown
        assert "project/bin/tool.exe: binary" in markdown
        assert "project/vendor.zip: nested archive" in markdown
        assert body["json"]["entries"] == 4
        assert body["json"]["text_files"] == 2
    finally:
        await client.close()

    context = await runtime_api._ensure_chat_attachment_context(session_id="s1", attachment_ids=[file_id])
    assert context["context_file_ids"] == [file_id]
    assert context["failures"] == []


@pytest.mark.asyncio
async def test_attachment_bytes_outlive_the_run_and_go_with_the_session(isolated_attachment_storage, monkeypatch):
    """The transcript's chip must still open the file after the run; deleting the session drops it."""
    from types import SimpleNamespace

    from src.gateway import runtime_api

    async def _delete_session(_session_id):
        return False  # a session that received uploads but was never persisted

    monkeypatch.setattr(runtime_api, "session_manager", SimpleNamespace(_initialized=True, delete_session=_delete_session))

    client = await _client()
    try:
        upload = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("2026-09 日志 (final).log", "错误：连接超时\n".encode("utf-8"), "text/plain"),
        )
        assert upload.status == 201, await upload.text()
        uploaded = await upload.json()
        assert uploaded["filename"] == "2026-09 日志 (final).log"
        file_id = uploaded["file_id"]

        context = await runtime_api._ensure_chat_attachment_context(session_id="s1", attachment_ids=[file_id])
        assert context["context_file_ids"] == [file_id]

        # What the chat handlers do once the run is over: release the context...
        await runtime_api._cleanup_one_shot_attachments("s1", [file_id])
        assert runtime_api.file_context_storage.get_file_meta("s1", file_id) is None

        # ...but keep the bytes for the chip in the transcript.
        inline = await client.get(f"/api/files/{file_id}", params={"session_id": "s1"})
        assert inline.status == 200
        assert inline.headers["Content-Disposition"].startswith("inline;")
        assert "filename*=UTF-8''2026-09%20%E6%97%A5%E5%BF%97%20%28final%29.log" in inline.headers["Content-Disposition"]
        assert await inline.read() == "错误：连接超时\n".encode("utf-8")

        download = await client.get(f"/api/files/{file_id}", params={"session_id": "s1", "download": "1"})
        assert download.status == 200
        assert download.headers["Content-Disposition"].startswith("attachment;")

        # Deleting the session takes its attachments with it, even when the
        # session itself was never persisted.
        deleted = await client.delete("/api/sessions/s1")
        assert deleted.status == 404
        gone = await client.get(f"/api/files/{file_id}", params={"session_id": "s1"})
        assert gone.status == 404
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_markup_attachments_never_render_inline(isolated_attachment_storage, monkeypatch):
    """An uploaded page must not execute under the Portal's origin when a chip opens it."""
    monkeypatch.setenv("EFP_CHAT_UPLOAD_EXTENSIONS", "html,svg,txt")
    client = await _client()
    try:
        page = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("page.html", b"<html><script>alert(1)</script></html>", "text/html"),
        )
        assert page.status == 201, await page.text()
        page_id = (await page.json())["file_id"]
        assert (await page.json())["content_type"] == "text/html"

        served = await client.get(f"/api/files/{page_id}", params={"session_id": "s1"})
        assert served.status == 200
        assert served.headers["Content-Disposition"].startswith("attachment;")
        assert served.headers["X-Content-Type-Options"] == "nosniff"

        note = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("note.txt", b"plain", "text/plain"),
        )
        note_id = (await note.json())["file_id"]
        served_text = await client.get(f"/api/files/{note_id}", params={"session_id": "s1"})
        assert served_text.headers["Content-Disposition"].startswith("inline;")
        assert served_text.headers["X-Content-Type-Options"] == "nosniff"
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_retention_sweep_drops_old_uploads_only(isolated_attachment_storage, monkeypatch):
    from datetime import datetime, timedelta

    from src.gateway import runtime_api
    from src.utils.file_parser.storage import sweep_expired_files

    client = await _client()
    try:
        old = await client.post("/api/files/upload", params={"session_id": "s1"}, data=_file_form("old.txt", b"old", "text/plain"))
        fresh = await client.post("/api/files/upload", params={"session_id": "s1"}, data=_file_form("new.txt", b"new", "text/plain"))
        old_id = (await old.json())["file_id"]
        fresh_id = (await fresh.json())["file_id"]

        storage_namespace = runtime_api.get_metadata.__globals__
        storage_namespace["_file_metadata"][old_id].uploaded_at = (datetime.utcnow() - timedelta(days=40)).isoformat() + "Z"

        assert sweep_expired_files(0) == 0  # disabled
        assert sweep_expired_files(30) == 1
        assert (await client.get(f"/api/files/{old_id}", params={"session_id": "s1"})).status == 404
        assert (await client.get(f"/api/files/{fresh_id}", params={"session_id": "s1"})).status == 200

        # The startup hook reads the env; a non-positive value keeps everything.
        monkeypatch.setenv("EFP_CHAT_UPLOAD_RETENTION_DAYS", "0")
        assert runtime_api._upload_retention_days() == 0
        await runtime_api._sweep_expired_attachments_on_startup(None)
        assert (await client.get(f"/api/files/{fresh_id}", params={"session_id": "s1"})).status == 200
        monkeypatch.setenv("EFP_CHAT_UPLOAD_RETENTION_DAYS", "garbage")
        assert runtime_api._upload_retention_days() == 30
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_display_attachments_describe_each_file_for_the_transcript(isolated_attachment_storage, monkeypatch):
    from src.gateway import runtime_api

    monkeypatch.setenv("EFP_CHAT_UPLOAD_EXTENSIONS", "png,txt,log")
    client = await _client()
    try:
        text_upload = await client.post(
            "/api/files/upload", params={"session_id": "s1"}, data=_file_form("app.log", b"ERROR boom", "text/plain")
        )
        image_upload = await client.post(
            "/api/files/upload", params={"session_id": "s1"}, data=_file_form("shot.png", PNG_BYTES, "image/png")
        )
        text_id = (await text_upload.json())["file_id"]
        image_id = (await image_upload.json())["file_id"]
    finally:
        await client.close()

    context = await runtime_api._ensure_chat_attachment_context(session_id="s1", attachment_ids=[text_id, image_id])
    display = runtime_api._build_display_attachments([text_id, image_id, "missing-file"], context)

    assert [item["name"] for item in display] == ["app.log", "shot.png"]
    assert display[0] == {
        "file_id": text_id, "id": text_id, "name": "app.log", "filename": "app.log",
        "content_type": "text/plain", "mime": "text/plain", "size": 10, "type": "file", "parsed": True,
    }
    assert display[1]["type"] == "image"
    assert display[1]["parsed"] is True
    assert display[1]["content_type"] == "image/png"


def test_history_shows_the_members_words_and_files_not_the_expanded_prompt():
    from src.gateway import runtime_api

    stored = {
        "role": "user",
        "content": "Based on the following context, answer the user's question.\n\nContext: ...\n\nQuestion: can you help check ths logs?\n\nAnswer:",
        "metadata": {
            "author_type": "human",
            "original_user_message": "can you help check ths logs?",
            "display_attachments": [
                {"file_id": "f1", "id": "f1", "name": "app.log", "content_type": "text/plain", "size": 233, "type": "file", "parsed": True},
            ],
            "internal_model_content_hidden": True,
        },
    }

    shown = runtime_api._normalize_chat_history_message(
        stored, portal_user_id=None, portal_user_name=None, runtime_agent_id=None, runtime_agent_name=None
    )

    assert shown["display_content"] == "can you help check ths logs?"
    assert shown["attachments"] == stored["metadata"]["display_attachments"]
    assert shown["metadata"]["internal_model_content_hidden"] is True
    # The model-facing content is still there for anyone who needs it.
    assert shown["content"].startswith("Based on the following context")

    # A turn without attachment metadata is untouched.
    plain = runtime_api._normalize_chat_history_message(
        {"role": "user", "content": "hello", "metadata": {"author_type": "human"}},
        portal_user_id=None, portal_user_name=None, runtime_agent_id=None, runtime_agent_name=None,
    )
    assert "display_content" not in plain and "attachments" not in plain


@pytest.mark.asyncio
async def test_preview_and_delete_roundtrip(isolated_attachment_storage):
    client = await _client()
    try:
        upload = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("notes.txt", b"preview me please", "text/plain"),
        )
        assert upload.status == 201
        file_id = (await upload.json())["file_id"]

        preview = await client.get(f"/api/files/{file_id}/preview", params={"session_id": "s1", "max_chars": "7"})
        assert preview.status == 200, await preview.text()
        body = await preview.json()
        assert body["success"] is True
        assert body["preview"] == "preview"
        assert body["truncated"] is True

        deleted = await client.delete(f"/api/files/{file_id}", params={"session_id": "s1"})
        assert deleted.status == 200
        assert (await deleted.json()) == {"success": True}

        gone = await client.get(f"/api/files/{file_id}", params={"session_id": "s1"})
        assert gone.status == 404
        deleted_again = await client.delete(f"/api/files/{file_id}", params={"session_id": "s1"})
        assert deleted_again.status == 404
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_attachment_bound_to_a_session_is_invisible_to_other_sessions(isolated_attachment_storage):
    client = await _client()
    try:
        upload = await client.post(
            "/api/files/upload",
            params={"session_id": "owner"},
            data=_file_form("notes.txt", b"private", "text/plain"),
        )
        file_id = (await upload.json())["file_id"]

        for session in ("intruder", None):
            params = {"session_id": session} if session else {}
            parsed = await client.post("/api/files/parse", params=params, json={"file_id": file_id})
            assert parsed.status == 404, session
            preview = await client.get(f"/api/files/{file_id}/preview", params=params)
            assert preview.status == 404, session
            raw = await client.get(f"/api/files/{file_id}", params=params)
            assert raw.status == 404, session
            deleted = await client.delete(f"/api/files/{file_id}", params=params)
            assert deleted.status == 404, session

        # The session id may also travel in the JSON body or the header.
        parsed = await client.post("/api/files/parse", json={"file_id": file_id, "session_id": "owner"})
        assert parsed.status == 200, await parsed.text()
        preview = await client.get(f"/api/files/{file_id}/preview", headers={"X-Session-ID": "owner"})
        assert preview.status == 200
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_upload_rejects_extension_not_on_the_allowlist(isolated_attachment_storage):
    client = await _client()
    try:
        response = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("tool.exe", b"MZ\x90\x00binary", "application/octet-stream"),
        )
        assert response.status == 415
        body = await response.json()
        assert body["success"] is False
        assert "File type .exe is not allowed" in body["error"]
        assert "Allowed: pdf, docx, xlsx, csv, txt, log, pptx, zip, md, yaml, yml, json, xml" in body["error"]

        # Images are off the default list: the default model has no vision.
        response = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("shot.png", PNG_BYTES, "image/png"),
        )
        assert response.status == 415
        assert "File type .png is not allowed" in (await response.json())["error"]

        response = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("README", b"no extension", "text/plain"),
        )
        assert response.status == 415
        assert "without an extension" in (await response.json())["error"]
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_upload_rejects_binary_bytes_behind_a_text_extension(isolated_attachment_storage):
    client = await _client()
    try:
        response = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("notes.txt", b"MZ\x90\x00\xff\xfe\x00\x01\x02\x03", "text/plain"),
        )
        assert response.status == 415
        body = await response.json()
        assert body["success"] is False
        assert "not a supported .txt file" in body["error"]
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_upload_allowlist_follows_efp_chat_upload_extensions(isolated_attachment_storage, monkeypatch):
    monkeypatch.setenv("EFP_CHAT_UPLOAD_EXTENSIONS", ".MD, json,txt")
    client = await _client()
    try:
        markdown = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("README.md", "# Title\n\n中文 paragraph\n".encode("utf-8"), "text/markdown"),
        )
        assert markdown.status == 201, await markdown.text()
        assert (await markdown.json())["content_type"] == "text/markdown"

        data = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("config.json", b'{"a": 1}', "application/json"),
        )
        assert data.status == 201, await data.text()
        json_upload = await data.json()
        assert json_upload["content_type"] == "application/json"

        parsed = await client.post("/api/files/parse", params={"session_id": "s1"}, json={"file_id": json_upload["file_id"]})
        assert parsed.status == 200, await parsed.text()
        assert '"a": 1' in (await parsed.json())["markdown"]

        # pdf is a default but not on this deployment's list.
        pdf = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("doc.pdf", b"%PDF-1.4\n%fake", "application/pdf"),
        )
        assert pdf.status == 415
        assert "Allowed: md, json, txt" in (await pdf.json())["error"]
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_upload_rejects_files_over_efp_max_upload_mb(isolated_attachment_storage, monkeypatch):
    monkeypatch.setenv("EFP_MAX_UPLOAD_MB", "1")
    client = await _client(client_max_size=8 * 1024 * 1024)
    try:
        response = await client.post(
            "/api/files/upload",
            params={"session_id": "s1"},
            data=_file_form("big.txt", b"a" * (1024 * 1024 + 16), "text/plain"),
        )
        assert response.status == 413
        body = await response.json()
        assert body["success"] is False
        assert "1MB" in body["error"]
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_upload_without_a_file_part_is_a_client_error(isolated_attachment_storage):
    client = await _client()
    try:
        # Multipart body with fields but no file part.
        form = FormData()
        form.add_field("session_id", "s1", content_type="text/plain")
        response = await client.post("/api/files/upload", data=form)
        assert response.status == 400
        assert (await response.json())["error"] == "No file provided"

        # Not multipart at all.
        for payload in ({"json": {"file": "not multipart"}}, {"data": {"session_id": "s1"}}):
            response = await client.post("/api/files/upload", **payload)
            assert response.status == 400
            assert (await response.json())["error"] == "multipart/form-data required"
    finally:
        await client.close()


@pytest.mark.asyncio
async def test_parse_validates_its_input(isolated_attachment_storage):
    client = await _client()
    try:
        response = await client.post("/api/files/parse", params={"session_id": "s1"}, json={})
        assert response.status == 400
        assert (await response.json())["error"] == "file_id is required"

        response = await client.post("/api/files/parse", params={"session_id": "s1"}, json={"file_id": "missing"})
        assert response.status == 404

        response = await client.post(
            "/api/files/parse",
            params={"session_id": "s1"},
            data=b"not json",
            headers={"Content-Type": "application/json"},
        )
        assert response.status == 400
        assert (await response.json())["error"] == "Invalid JSON body"
    finally:
        await client.close()
